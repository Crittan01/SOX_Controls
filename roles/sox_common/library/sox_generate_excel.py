#!/usr/bin/python
# -*- coding: utf-8 -*-

# Copyright: (c) 2025, NTT Data
# GNU General Public License v3.0+ (see COPYING or https://www.gnu.org/licenses/gpl-3.0.txt)

from __future__ import absolute_import, division, print_function
__metaclass__ = type

DOCUMENTATION = r'''
---
module: sox_generate_excel
short_description: Genera archivos Excel para reportes SOX
version_added: "1.0.0"
description:
    - Genera archivos Excel formateados con datos de controles SOX
    - Soporta headers personalizados, metadata y estilos
    - Utiliza openpyxl para generar archivos .xlsx
options:
    output_file:
        description: Ruta completa del archivo Excel a generar
        required: true
        type: str
    sheet_name:
        description: Nombre de la hoja de Excel
        required: false
        type: str
        default: "Datos"
    headers:
        description: Lista de headers para las columnas
        required: true
        type: list
        elements: str
    data:
        description: Lista de diccionarios con los datos a insertar
        required: true
        type: list
        elements: dict
    data_keys:
        description: Lista de claves de los diccionarios en el orden de las columnas
        required: true
        type: list
        elements: str
    title:
        description: Título del reporte (se coloca en la primera fila)
        required: false
        type: str
    metadata:
        description: Lista de líneas de metadata para agregar antes de los datos
        required: false
        type: list
        elements: str
        default: []
author:
    - Equipo Automatización NTT Data
'''

EXAMPLES = r'''
- name: Generar Excel con usuarios SOX
  sox_generate_excel:
    output_file: "/tmp/sox46_usuarios.xlsx"
    sheet_name: "Usuarios Creados"
    headers:
      - "Base de Datos"
      - "Usuario"
      - "Estado"
    data: "{{ usuarios_estructurados }}"
    data_keys:
      - "base_de_datos"
      - "usuario"
      - "estado"
    title: "SOX-046 - Usuarios Creados"
    metadata:
      - "Fecha: 2025-01-05"
      - "Ejecutado por: ansible"
'''

RETURN = r'''
file_path:
    description: Ruta completa del archivo generado
    type: str
    returned: success
    sample: "/tmp/sox46_usuarios_2025-01-05.xlsx"
rows_written:
    description: Número de filas de datos escritas
    type: int
    returned: success
    sample: 150
message:
    description: Mensaje de resultado
    type: str
    returned: always
    sample: "Excel generado exitosamente con 150 registros"
'''

import os
import traceback
from ansible.module_utils.basic import AnsibleModule

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def create_excel_report(module):
    """
    Crea un archivo Excel con los datos proporcionados
    """
    output_file = module.params['output_file']
    sheet_name = module.params['sheet_name']
    headers = module.params['headers']
    data = module.params['data']
    data_keys = module.params['data_keys']
    title = module.params.get('title')
    metadata = module.params.get('metadata', [])

    try:
        # Crear workbook
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name[:31]  # Excel limita nombres a 31 caracteres

        current_row = 1

        # Estilos
        title_font = Font(name='Calibri', size=16, bold=True, color='FFFFFF')
        title_fill = PatternFill(start_color='0066CC', end_color='0066CC', fill_type='solid')
        header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
        metadata_font = Font(name='Calibri', size=10, italic=True)
        data_font = Font(name='Calibri', size=10)
        center_align = Alignment(horizontal='center', vertical='center')
        left_align = Alignment(horizontal='left', vertical='center')
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Agregar título si existe
        if title:
            ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
            title_cell = ws.cell(row=current_row, column=1)
            title_cell.value = title
            title_cell.font = title_font
            title_cell.fill = title_fill
            title_cell.alignment = center_align
            ws.row_dimensions[current_row].height = 25
            current_row += 1

        # Agregar metadata
        if metadata:
            for meta_line in metadata:
                ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=len(headers))
                meta_cell = ws.cell(row=current_row, column=1)
                meta_cell.value = meta_line
                meta_cell.font = metadata_font
                meta_cell.alignment = left_align
                current_row += 1
            current_row += 1  # Espacio en blanco

        # Agregar headers
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = center_align
            cell.border = border
        ws.row_dimensions[current_row].height = 20
        current_row += 1

        # Agregar datos
        rows_written = 0
        for record in data:
            for col_num, key in enumerate(data_keys, 1):
                value = record.get(key, '')
                # Convertir None o 'None' a vacío
                if value is None or value == 'None' or value == 'N/A':
                    value = ''
                cell = ws.cell(row=current_row, column=col_num)
                cell.value = value
                cell.font = data_font
                cell.alignment = center_align if col_num > 1 else left_align
                cell.border = border
            current_row += 1
            rows_written += 1

        # Ajustar ancho de columnas
        for col_num, header in enumerate(headers, 1):
            # Calcular ancho basado en el header
            max_length = len(str(header))
            # Revisar los datos
            for record in data[:100]:  # Solo revisar primeros 100 para performance
                key = data_keys[col_num - 1]
                value = str(record.get(key, ''))
                if len(value) > max_length:
                    max_length = len(value)
            # Establecer ancho (máximo 50 caracteres)
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[chr(64 + col_num)].width = adjusted_width

        # Guardar archivo
        wb.save(output_file)

        return {
            'changed': True,
            'file_path': output_file,
            'rows_written': rows_written,
            'message': f'Excel generado exitosamente con {rows_written} registros en {output_file}'
        }

    except Exception as e:
        module.fail_json(
            msg=f'Error al generar archivo Excel: {str(e)}',
            exception=traceback.format_exc()
        )


def main():
    module = AnsibleModule(
        argument_spec=dict(
            output_file=dict(type='str', required=True),
            sheet_name=dict(type='str', default='Datos'),
            headers=dict(type='list', elements='str', required=True),
            data=dict(type='list', elements='dict', required=True),
            data_keys=dict(type='list', elements='str', required=True),
            title=dict(type='str', default=None),
            metadata=dict(type='list', elements='str', default=[])
        ),
        supports_check_mode=False
    )

    if not HAS_OPENPYXL:
        module.fail_json(msg='El módulo openpyxl es requerido. Instale con: pip install openpyxl')

    # Validar que headers y data_keys tengan la misma longitud
    if len(module.params['headers']) != len(module.params['data_keys']):
        module.fail_json(msg='La cantidad de headers debe coincidir con la cantidad de data_keys')

    # Crear directorio de salida si no existe
    output_dir = os.path.dirname(module.params['output_file'])
    if output_dir and not os.path.exists(output_dir):
        try:
            os.makedirs(output_dir, mode=0o755)
        except Exception as e:
            module.fail_json(msg=f'No se pudo crear el directorio {output_dir}: {str(e)}')

    result = create_excel_report(module)
    module.exit_json(**result)


if __name__ == '__main__':
    main()
